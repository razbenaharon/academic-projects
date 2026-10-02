from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "000000000_000000000.docx"

# Base: compact_reference_guide. Named one-page Hebrew submission overrides:
# Arial throughout; 0.75-inch margins; 10.5-point body; 1.05 spacing; no
# running header/footer. These overrides keep the required explanation to one
# page while preserving a clear hierarchy and comfortable Hebrew rendering.
NAVY = RGBColor(31, 78, 121)
INK = RGBColor(24, 24, 24)
MUTED = RGBColor(90, 90, 90)


def set_cell_or_paragraph_bidi(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    bidi = p_pr.find(qn("w:bidi"))
    if bidi is None:
        bidi = OxmlElement("w:bidi")
        p_pr.append(bidi)
    bidi.set(qn("w:val"), "1")


def set_run_font(run, size, *, bold=False, color=INK):
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = color
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        r_fonts.set(qn(f"w:{attr}"), "Arial")
    size_cs = r_pr.find(qn("w:szCs"))
    if size_cs is None:
        size_cs = OxmlElement("w:szCs")
        r_pr.append(size_cs)
    size_cs.set(qn("w:val"), str(int(size * 2)))
    rtl = r_pr.find(qn("w:rtl"))
    if rtl is None:
        rtl = OxmlElement("w:rtl")
        r_pr.append(rtl)
    rtl.set(qn("w:val"), "1")


def configure_style(style, size, *, bold=False, color=INK, before=0, after=3):
    style.font.name = "Arial"
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = color
    style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    style._element.rPr.rFonts.set(qn("w:cs"), "Arial")
    size_cs = style._element.rPr.find(qn("w:szCs"))
    if size_cs is None:
        size_cs = OxmlElement("w:szCs")
        style._element.rPr.append(size_cs)
    size_cs.set(qn("w:val"), str(int(size * 2)))
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.line_spacing = 1.05
    style.paragraph_format.keep_with_next = style.name.startswith("Heading")


def add_rtl_paragraph(doc, text, *, style=None, size=10.5, bold=False,
                      color=INK, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
                      before=0, after=3, keep_with_next=False):
    paragraph = doc.add_paragraph(style=style)
    set_cell_or_paragraph_bidi(paragraph)
    paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.05
    paragraph.paragraph_format.keep_with_next = keep_with_next
    run = paragraph.add_run(text)
    set_run_font(run, size, bold=bold, color=color)
    return paragraph


def add_heading(doc, text):
    return add_rtl_paragraph(
        doc,
        text,
        style="Heading 1",
        size=13,
        bold=True,
        color=NAVY,
        align=WD_ALIGN_PARAGRAPH.RIGHT,
        before=7,
        after=3,
        keep_with_next=True,
    )


def build_document():
    doc = Document()
    section = doc.sections[0]
    section.start_type = WD_SECTION.NEW_PAGE
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)

    configure_style(doc.styles["Normal"], 10.5, after=4)
    configure_style(doc.styles["Title"], 18, bold=True, color=NAVY,
                    before=0, after=2)
    title_p_pr = doc.styles["Title"]._element.get_or_add_pPr()
    title_border = title_p_pr.find(qn("w:pBdr"))
    if title_border is not None:
        title_p_pr.remove(title_border)
    configure_style(doc.styles["Heading 1"], 13, bold=True, color=NAVY,
                    before=7, after=3)
    configure_style(doc.styles["Heading 2"], 11.5, bold=True, color=NAVY,
                    before=4, after=2)
    configure_style(doc.styles["Heading 3"], 10.5, bold=True, color=NAVY,
                    before=3, after=2)

    title = add_rtl_paragraph(
        doc,
        "תרגיל בית 3 - אסטרטגיות הצעה אדפטיביות במכרז GSP",
        style="Title",
        size=18,
        bold=True,
        color=NAVY,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        after=2,
        keep_with_next=True,
    )
    title.paragraph_format.line_spacing = 1.0

    add_rtl_paragraph(
        doc,
        "מזהי מגישים: 000000000, 000000000  |  קובץ מימוש: id_000000000_000000000.py",
        size=9.5,
        color=MUTED,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        after=5,
        keep_with_next=True,
    )

    add_heading(doc, "הרעיון הכללי ולמידת השוק")
    add_rtl_paragraph(
        doc,
        "בכל סימולציה נערכים מכרזי GSP חוזרים. סוכן 1 ממקסם תועלת ללא מגבלת "
        "תקציב, וסוכן 2 ממקסם תועלת תחת תקציב כולל, כאשר ליתרה בסוף הסימולציה "
        "אין ערך. אחרי כל סיבוב מתפרסמות תוצאות הזוכים בפורמט (מזהה, משבצת, "
        "מחיר). במכרז GSP המחיר של משבצת שווה להצעה של הסוכן המדורג מיד אחריה; "
        "לכן ניתן לשחזר במדויק את הצעות כל הזוכים שאינם במקום הראשון. להצעת "
        "המוביל נשמר אומדן שמרני המבוסס על חסם תחתון מן המחיר שנצפה וממידע קודם.",
    )
    add_rtl_paragraph(
        doc,
        "שני הסוכנים שומרים חלון מתגלגל וחסום של צילומי הצעות היריבים, ממוינים "
        "בסדר יורד. הצעות מועמדות נבנות מרשת בקפיצות של חמישה אחוזים מן הערך "
        "הפרטי ומסכומים זעירים "
        "מעל גבולות דירוג שנצפו בחמשת הסיבובים האחרונים. עבור כל מועמד מתבצע "
        "שחזור נגד-עובדתי על 30 הצילומים האחרונים, שממנו נאמדים המשבצת, המחיר, "
        "התועלת הצפויה וההוצאה הצפויה. החלונות הקבועים מאפשרים הסתגלות מהירה "
        "ושומרים את זמן הריצה והזיכרון בלתי תלויים במספר הסיבובים הכולל.",
    )

    add_heading(doc, "סוכן 1 - ללא מגבלת תקציב")
    add_rtl_paragraph(
        doc,
        "בסיבוב הראשון הסוכן מציע את מלוא הערך כדי לקבל תצפית שוק אינפורמטיבית. "
        "לאחר מכן הוא בוחר את ההצעה בעלת התועלת האמפירית הגבוהה ביותר - מכפלת "
        "שיעור ההקלקה בהפרש שבין הערך למחיר; במקרה של תיקו נבחרת ההצעה הנמוכה "
        "יותר. רצפת הצעה בגובה 65 אחוזים מן הערך מונעת ויתור מופרז על דירוג מול המתחרים, אך "
        "הלומד עדיין רשאי להציע יותר כאשר משבצת גבוהה מצדיקה זאת. ההצעה לעולם "
        "אינה עולה על הערך הפרטי.",
    )

    add_heading(doc, "סוכן 2 - ניהול תקציב")
    add_rtl_paragraph(
        doc,
        "לכל הצעה מועמדת נאמדות תועלת והוצאה. יעד ההוצאה לסיבוב הוא 95 אחוזים "
        "מן התקציב שנותר, חלקי מספר הסיבובים שנותרו. מקדם זה יוצר רזרבה קטנה "
        "כנגד שגיאת חיזוי, והחישוב מחדש בכל סיבוב משחרר יתרה שלא נוצלה בהמשך. "
        "הסוכן משתמש במחיר צל לתקציב ובוחר פעולה הממקסמת את התועלת הצפויה פחות "
        "מחיר הצל כפול ההוצאה הצפויה. חיפוש בינארי חסום של 20 "
        "איטרציות מאתר שתי פעולות יעילות משני צדי יעד ההוצאה, והסוכן מערבב "
        "ביניהן בהסתברות המתאימה. כל הצעה מוגבלת לערך הפרטי ולתקציב שנותר.",
    )

    add_heading(doc, "בדיקות, חלופות ובטיחות")
    add_rtl_paragraph(
        doc,
        "בהשוואות מקומיות נעשה שימוש באותם זרעים ובהחלפת הערכים והתקציבים בין "
        "הסוכנים. נבדקו אורכי היסטוריה שונים, שקלול אקספוננציאלי, הצעות קבועות "
        "בשיעורים שבין 55 ל-85 אחוזים, יעדי קצב הוצאה שבין 80 ל-125 אחוזים ורצפות "
        "הצעה תלויות תקציב. "
        "השילוב שנבחר היה היציב ביותר בניסויים. המימוש משתמש רק בספרייה "
        "הסטנדרטית, מחזיר הצעות סופיות ולא שליליות, משתמש במבנים חסומי גודל "
        "ובחיפושים חסומים, ונועד לעמוד במגבלת 50 מילישניות לכל קריאה.",
        after=0,
    )

    props = doc.core_properties
    props.title = "תרגיל בית 3 - הסבר על אסטרטגיות ההצעה"
    props.subject = "הסבר בעברית למימוש BiddingAgent1 ו-BiddingAgent2"
    props.author = ""
    props.last_modified_by = ""
    props.keywords = "GSP, bidding, pacing, Hebrew"
    doc.save(OUTPUT)


if __name__ == "__main__":
    build_document()
    print(OUTPUT)
